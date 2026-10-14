"""Liquidaciones y Fondos — registro de lo que se le debe y se le pagó a cada persona.
Ejecutar:  python app.py   y abrir http://127.0.0.1:5002
"""
import io
import os
import threading
import uuid
import webbrowser
from collections import defaultdict
from datetime import date, datetime

from flask import Flask, abort, flash, redirect, render_template, request, send_file, send_from_directory, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename

import db

VERSION = "1.0.0"
PUERTO = 5002
EXTENSIONES = {".pdf", ".png", ".jpg", ".jpeg", ".webp"}

db.init()
if db.vacia() and os.path.exists(os.path.join(db.BASE, "datos_iniciales.xlsx")):
    import importar_excel
    p, i = importar_excel.importar(os.path.join(db.BASE, "datos_iniciales.xlsx"))
    print(f"Primera vez: importadas {p} personas y {i} conceptos desde datos_iniciales.xlsx")

app = Flask(__name__)
app.secret_key = "liquidaciones-y-fondos-local"
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024


# ---------------------------------------------------------------- formato
@app.template_filter("pesos")
def pesos(v):
    if v is None:
        return ""
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return "$ " + s


@app.template_filter("fecha")
def fecha(v):
    if not v:
        return ""
    try:
        return datetime.strptime(v[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return v


@app.context_processor
def globales():
    return {"TIPOS": db.TIPOS, "TIPOS_CORTO": db.TIPOS_CORTO, "TIPOS_NOMBRE": db.TIPOS_NOMBRE,
            "version": VERSION, "hoy": date.today().isoformat()}


def _monto(txt):
    """Acepta 1234567,89 / 1.234.567,89 / 1234567.89 / $ 1.234."""
    txt = (txt or "").replace("$", "").replace(" ", "").strip()
    if not txt:
        return None
    if "," in txt:
        txt = txt.replace(".", "").replace(",", ".")
    elif txt.count(".") > 1 or (txt.count(".") == 1 and len(txt.split(".")[1]) == 3):
        txt = txt.replace(".", "")
    try:
        return round(float(txt), 2)
    except ValueError:
        return None


def _volver(defecto="pagos"):
    sig = request.form.get("next") or request.args.get("next")
    if sig and sig.startswith("/") and not sig.startswith("//"):
        return redirect(sig)
    return redirect(url_for(defecto))


# ---------------------------------------------------------------- datos
def personas_con_items(q="", obra="", estado=""):
    with db.conn() as c:
        personas = [dict(r) for r in c.execute("SELECT * FROM personal ORDER BY nombre")]
        items = [dict(r) for r in c.execute("SELECT * FROM items ORDER BY tipo, id")]
    por_persona = defaultdict(list)
    for it in items:
        por_persona[it["personal_id"]].append(it)
    salida = []
    q = q.strip().upper()
    for p in personas:
        its = por_persona[p["id"]]
        p["items"] = {t: [i for i in its if i["tipo"] == t] for t, _ in db.TIPOS}
        p["lista_items"] = its
        p["resumen"] = {}
        for t, _ in db.TIPOS:
            ti = p["items"][t]
            pend = [i for i in ti if i["estado"] == "PENDIENTE"]
            p["resumen"][t] = {
                "n": len(ti), "n_pend": len(pend),
                "deuda": sum(i["monto"] or 0 for i in pend),
                "pagado": sum(i["monto"] or 0 for i in ti if i["estado"] == "PAGADO"),
                "sin_monto": sum(1 for i in pend if i["monto"] is None),
            }
        p["n_comp"] = sum(1 for i in its if i["comprobante"])
        if not its:
            p["estado_total"] = "SIN CARGAR"
        elif any(i["estado"] == "PENDIENTE" for i in its):
            p["estado_total"] = "NO PAGADO"
        else:
            p["estado_total"] = "PAGADO"
        p["deuda"] = sum(i["monto"] or 0 for i in its if i["estado"] == "PENDIENTE")
        p["pagado"] = sum(i["monto"] or 0 for i in its if i["estado"] == "PAGADO")
        if q and q not in (p["nombre"] or "").upper() and q not in (p["cuil"] or ""):
            continue
        if obra and (p["obra"] or "") != obra:
            continue
        if estado and p["estado_total"] != estado:
            continue
        salida.append(p)
    return salida


def obras():
    with db.conn() as c:
        return [r[0] for r in c.execute("SELECT DISTINCT obra FROM personal WHERE obra IS NOT NULL AND obra<>'' ORDER BY obra")]


# ---------------------------------------------------------------- pestañas
@app.route("/")
def resumen():
    with db.conn() as c:
        items = [dict(r) for r in c.execute(
            "SELECT i.*, p.obra FROM items i JOIN personal p ON p.id=i.personal_id")]
    tot = {"deuda": 0, "pagado": 0, "n_pend": 0, "n_pag": 0, "sin_monto": 0}
    por_tipo = {t: {"deuda": 0, "pagado": 0, "n_pend": 0, "n_pag": 0, "sin_monto": 0} for t, _ in db.TIPOS}
    por_obra = defaultdict(lambda: {"deuda": 0, "pagado": 0, "n_pend": 0})
    for it in items:
        m = it["monto"] or 0
        dest = [tot, por_tipo.setdefault(it["tipo"], {"deuda": 0, "pagado": 0, "n_pend": 0, "n_pag": 0, "sin_monto": 0})]
        for d in dest:
            if it["estado"] == "PAGADO":
                d["pagado"] += m
                d["n_pag"] += 1
            else:
                d["deuda"] += m
                d["n_pend"] += 1
                if it["monto"] is None:
                    d["sin_monto"] += 1
        o = por_obra[it["obra"] or "Sin obra"]
        if it["estado"] == "PAGADO":
            o["pagado"] += m
        else:
            o["deuda"] += m
            o["n_pend"] += 1
    personas = personas_con_items()
    cuenta = defaultdict(int)
    for p in personas:
        cuenta[p["estado_total"]] += 1
    obras_orden = sorted(por_obra.items(), key=lambda kv: -kv[1]["deuda"])
    return render_template("resumen.html", tot=tot, por_tipo=por_tipo, por_obra=obras_orden,
                           cuenta=cuenta, n_personas=len(personas),
                           top=sorted([p for p in personas if p["deuda"]], key=lambda p: -p["deuda"])[:10])


@app.route("/pagos")
def pagos():
    q, obra, estado = request.args.get("q", ""), request.args.get("obra", ""), request.args.get("estado", "")
    personas = personas_con_items(q, obra, estado)
    return render_template("pagos.html", personas=personas, obras=obras(), q=q, obra=obra, estado=estado)


@app.route("/cargar")
def cargar():
    tipo = request.args.get("tipo", "")
    sin_monto = request.args.get("sin_monto") == "1"
    pid = request.args.get("persona", "")
    sql = ("SELECT i.*, p.nombre, p.cuil, p.obra FROM items i JOIN personal p ON p.id=i.personal_id WHERE 1=1")
    args = []
    if tipo:
        sql += " AND i.tipo=?"
        args.append(tipo)
    if sin_monto:
        sql += " AND i.monto IS NULL"
    if pid:
        sql += " AND i.personal_id=?"
        args.append(pid)
    sql += " ORDER BY i.id DESC LIMIT 400"
    with db.conn() as c:
        items = [dict(r) for r in c.execute(sql, args)]
        personas = [dict(r) for r in c.execute("SELECT id, nombre, cuil FROM personal ORDER BY nombre")]
    return render_template("cargar.html", items=items, personas=personas, tipo=tipo, sin_monto=sin_monto, persona=pid)


@app.route("/personal")
def personal():
    q = request.args.get("q", "")
    return render_template("personal.html", personas=personas_con_items(q), q=q, obras=obras(), editar=None)


def _conceptos_elegidos():
    elegidos = [t for t in request.args.getlist("c") if t in db.TIPOS_NOMBRE]
    return elegidos or [t for t, _ in db.TIPOS]


@app.route("/personal/<int:pid>")
def ficha(pid):
    p = next((x for x in personas_con_items() if x["id"] == pid), None)
    if not p:
        abort(404)
    elegidos = _conceptos_elegidos()
    items = [i for i in p["lista_items"] if i["tipo"] in elegidos]
    items.sort(key=lambda i: ([t for t, _ in db.TIPOS].index(i["tipo"]), i["id"]))
    deuda = sum(i["monto"] or 0 for i in items if i["estado"] == "PENDIENTE")
    pagado = sum(i["monto"] or 0 for i in items if i["estado"] == "PAGADO")
    return render_template("ficha.html", p=p, obras=obras(), elegidos=elegidos, items=items,
                           deuda=deuda, pagado=pagado, ahora=datetime.now())


@app.route("/deuda")
def deuda():
    elegidos = _conceptos_elegidos()
    q, obra = request.args.get("q", ""), request.args.get("obra", "")
    solo_deuda = request.args.get("todos") != "1"
    filas = []
    for p in personas_con_items(q, obra):
        r = {t: p["resumen"][t] for t in elegidos}
        total = sum(x["deuda"] for x in r.values())
        pend = sum(x["n_pend"] for x in r.values())
        if solo_deuda and not pend:
            continue
        filas.append({"p": p, "r": r, "total": total, "pend": pend})
    return render_template("deuda.html", filas=filas, elegidos=elegidos, obras=obras(), q=q, obra=obra,
                           solo_deuda=solo_deuda, ahora=datetime.now())


@app.post("/personal/guardar")
def personal_guardar():
    f = request.form
    pid = f.get("id")
    nombre = " ".join(f.get("nombre", "").upper().split())
    cuil = "".join(ch for ch in f.get("cuil", "") if ch.isdigit()) or None
    if not nombre:
        flash("Falta el apellido y nombre.", "error")
        return _volver("personal")
    datos = (cuil, nombre, f.get("fecha_baja") or None, (" ".join(f.get("obra", "").upper().split()) or None),
             f.get("cuenta_sueldo", "").strip() or None, f.get("cuenta_fondo", "").strip() or None,
             f.get("telegrama") or None, f.get("notas", "").strip() or None)
    with db.conn() as c:
        otro = c.execute("SELECT id, nombre FROM personal WHERE cuil=? AND id<>?", (cuil, pid or 0)).fetchone() if cuil else None
        if otro:
            flash(f"Ya existe una persona con ese CUIL: {otro['nombre']}.", "error")
            return _volver("personal")
        if pid:
            c.execute("UPDATE personal SET cuil=?, nombre=?, fecha_baja=?, obra=?, cuenta_sueldo=?, cuenta_fondo=?,"
                      " telegrama=?, notas=? WHERE id=?", datos + (pid,))
            flash("Datos actualizados.", "ok")
        else:
            pid = c.execute("INSERT INTO personal (cuil, nombre, fecha_baja, obra, cuenta_sueldo, cuenta_fondo,"
                            " telegrama, notas) VALUES (?,?,?,?,?,?,?,?)", datos).lastrowid
            flash(f"{nombre} registrado. Ahora podés cargarle liquidación, FCL, aguinaldo o bonos.", "ok")
    return redirect(url_for("ficha", pid=pid))


@app.post("/personal/<int:pid>/borrar")
def personal_borrar(pid):
    with db.conn() as c:
        for (arch,) in c.execute("SELECT comprobante FROM items WHERE personal_id=? AND comprobante IS NOT NULL", (pid,)):
            _borrar_archivo(arch)
        c.execute("DELETE FROM personal WHERE id=?", (pid,))
    flash("Persona eliminada con todos sus conceptos.", "ok")
    return redirect(url_for("personal"))


# ---------------------------------------------------------------- conceptos (items)
def _borrar_archivo(nombre):
    if nombre:
        ruta = os.path.join(db.COMPROBANTES_DIR, nombre)
        if os.path.exists(ruta):
            os.remove(ruta)


def _guardar_comprobante(archivo, pid, tipo):
    ext = os.path.splitext(archivo.filename)[1].lower()
    if ext not in EXTENSIONES:
        flash("El comprobante tiene que ser PDF o imagen (JPG, PNG, WEBP).", "error")
        return None, None
    nombre = f"{pid}_{tipo}_{date.today():%Y%m%d}_{uuid.uuid4().hex[:6]}{ext}"
    archivo.save(os.path.join(db.COMPROBANTES_DIR, nombre))
    return nombre, secure_filename(archivo.filename) or nombre


@app.post("/item/guardar")
def item_guardar():
    f = request.form
    iid, pid, tipo = f.get("id"), f.get("personal_id"), f.get("tipo")
    if tipo not in db.TIPOS_NOMBRE:
        abort(400)
    estado = "PAGADO" if f.get("estado") == "PAGADO" else "PENDIENTE"
    monto = _monto(f.get("monto"))
    if f.get("monto", "").strip() and monto is None:
        flash("No se entendió el monto, revisalo.", "error")
        return _volver()
    fecha_pago = f.get("fecha_pago") or None
    if estado == "PENDIENTE":
        fecha_pago = None
    with db.conn() as c:
        if iid:
            viejo = c.execute("SELECT * FROM items WHERE id=?", (iid,)).fetchone()
            if not viejo:
                abort(404)
            pid = viejo["personal_id"]
        archivo = request.files.get("comprobante")
        comp = comp_nombre = None
        if archivo and archivo.filename:
            comp, comp_nombre = _guardar_comprobante(archivo, pid, tipo)
            if not comp:
                return _volver()
        if iid:
            c.execute("UPDATE items SET tipo=?, periodo=?, monto=?, estado=?, fecha_pago=?, notas=? WHERE id=?",
                      (tipo, f.get("periodo", "").strip() or None, monto, estado, fecha_pago,
                       f.get("notas", "").strip() or None, iid))
            if comp:
                _borrar_archivo(viejo["comprobante"])
                c.execute("UPDATE items SET comprobante=?, comprobante_nombre=? WHERE id=?", (comp, comp_nombre, iid))
        else:
            c.execute("INSERT INTO items (personal_id, tipo, periodo, monto, estado, fecha_pago, notas, comprobante,"
                      " comprobante_nombre) VALUES (?,?,?,?,?,?,?,?,?)",
                      (pid, tipo, f.get("periodo", "").strip() or None, monto, estado, fecha_pago,
                       f.get("notas", "").strip() or None, comp, comp_nombre))
    flash(f"{db.TIPOS_NOMBRE[tipo]} guardado" + (" con comprobante." if comp else "."), "ok")
    return _volver()


@app.post("/item/<int:iid>/monto")
def item_monto(iid):
    """Carga rápida de monto/período desde la pestaña Aguinaldos y fondos."""
    monto = _monto(request.form.get("monto"))
    with db.conn() as c:
        c.execute("UPDATE items SET monto=?, periodo=? WHERE id=?",
                  (monto, request.form.get("periodo", "").strip() or None, iid))
    flash("Monto guardado.", "ok")
    return _volver("cargar")


@app.post("/item/<int:iid>/borrar")
def item_borrar(iid):
    with db.conn() as c:
        r = c.execute("SELECT comprobante FROM items WHERE id=?", (iid,)).fetchone()
        if r:
            _borrar_archivo(r["comprobante"])
        c.execute("DELETE FROM items WHERE id=?", (iid,))
    flash("Concepto eliminado.", "ok")
    return _volver()


@app.post("/item/<int:iid>/quitar_comprobante")
def item_quitar_comprobante(iid):
    with db.conn() as c:
        r = c.execute("SELECT comprobante FROM items WHERE id=?", (iid,)).fetchone()
        if r:
            _borrar_archivo(r["comprobante"])
        c.execute("UPDATE items SET comprobante=NULL, comprobante_nombre=NULL WHERE id=?", (iid,))
    flash("Comprobante quitado.", "ok")
    return _volver()


@app.route("/comprobante/<int:iid>")
def comprobante(iid):
    with db.conn() as c:
        r = c.execute("SELECT comprobante, comprobante_nombre FROM items WHERE id=?", (iid,)).fetchone()
    if not r or not r["comprobante"]:
        abort(404)
    return send_from_directory(db.COMPROBANTES_DIR, r["comprobante"], download_name=r["comprobante_nombre"])


# ---------------------------------------------------------------- Excel
@app.route("/exportar")
def exportar():
    personas = personas_con_items(request.args.get("q", ""), request.args.get("obra", ""), request.args.get("estado", ""))
    wb = Workbook()
    ws = wb.active
    ws.title = "Liquidaciones y fondos"
    tit = ["CUIL", "APELLIDO Y NOMBRE", "FECHA DE BAJA", "LIQUIDACION FINAL", "FCL", "AGUINALDO", "BONOS ADUA",
           "ESTADO TOTAL", "OBRA", "CUENTA SUELDO", "CUENTA FONDO", "TELEGRAMA", "DEUDA ($)", "PAGADO ($)", "COMPROBANTES"]
    ws.append(tit)
    for i in range(1, len(tit) + 1):
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="1F4E78")

    def celda(its):
        partes = []
        for it in its:
            m = pesos(it["monto"]) if it["monto"] is not None else ""
            partes.append(f"{'PAGADO' if it['estado'] == 'PAGADO' else 'NO PAGADO'} {m}".strip())
        return " / ".join(partes)

    for p in personas:
        comps = [db.TIPOS_CORTO[i["tipo"]] for i in p["lista_items"] if i["comprobante"]]
        ws.append([p["cuil"], p["nombre"], fecha(p["fecha_baja"]), *(celda(p["items"][t]) for t, _ in db.TIPOS),
                   p["estado_total"], p["obra"], p["cuenta_sueldo"], p["cuenta_fondo"], p["telegrama"],
                   p["deuda"], p["pagado"], ", ".join(comps)])
    for i, w in enumerate([14, 34, 12, 22, 22, 22, 22, 13, 18, 18, 18, 11, 15, 15, 24], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for fila in ws.iter_rows(min_row=2, min_col=13, max_col=14):
        for cel in fila:
            cel.number_format = '"$" #,##0.00'
    ws.freeze_panes = "C2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"liquidaciones_y_fondos_{date.today():%Y-%m-%d}.xlsx")


if __name__ == "__main__":
    if os.environ.get("LF_NO_BROWSER") != "1":
        threading.Timer(1.2, lambda: webbrowser.open(f"http://127.0.0.1:{PUERTO}")).start()
    app.run(host="127.0.0.1", port=PUERTO, debug=False)
