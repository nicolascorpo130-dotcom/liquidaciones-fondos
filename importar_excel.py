"""Importa la planilla "liquidaciones y fondos.xlsx" a la base.

Se ejecuta solo la primera vez que se abre la app (si la base está vacía y existe
datos_iniciales.xlsx), o a mano:  python importar_excel.py archivo.xlsx

Criterio para cada celda de Liquidación / FCL / Aguinaldo / Bonos:
- vacía               -> no se crea el concepto (se carga más adelante)
- "PAGADO"            -> concepto pagado, monto sin cargar
- "NO PAGADO"         -> concepto pendiente, monto sin cargar
- un número           -> monto; pagado si la columna ESTADO TOTAL dice PAGADO, si no pendiente
- otro texto (FLF...) -> concepto pendiente, monto sin cargar, el texto queda como nota
Si COMPROBANTE LIQUIDACION tiene una fecha, la liquidación queda pagada en esa fecha.
"""
import sys
from datetime import date, datetime

from openpyxl import load_workbook

import db

COLUMNAS = {"LIQUIDACION": 3, "FCL": 4, "AGUINALDO": 5, "BONO": 6}  # índice 0 = columna A


def _fecha(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return None


def _cuenta(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return None if s in ("", "#N/A") else s


def _obra(v):
    if not v:
        return None
    s = " ".join(str(v).upper().split())
    if s.startswith("OBRA "):
        s = s[5:]
    return s


def importar(ruta):
    db.init()
    ws = load_workbook(ruta, data_only=True).active
    # La fila de títulos es la que tiene "CUIL" en la columna A
    fila_tit = next(r for r in range(1, 11) if str(ws.cell(r, 1).value or "").strip().upper() == "CUIL")
    personas = items = 0
    with db.conn() as c:
        for fila in ws.iter_rows(min_row=fila_tit + 1, values_only=True):
            fila = list(fila) + [None] * 14
            if not fila[1]:
                continue
            cuil = _cuenta(fila[0])
            estado_total = str(fila[7] or "").strip().upper()
            existe = c.execute("SELECT id FROM personal WHERE cuil=?", (cuil,)).fetchone() if cuil else None
            if existe:
                continue
            cur = c.execute(
                "INSERT INTO personal (cuil, nombre, fecha_baja, obra, cuenta_sueldo, cuenta_fondo, telegrama)"
                " VALUES (?,?,?,?,?,?,?)",
                (cuil, " ".join(str(fila[1]).upper().split()), _fecha(fila[2]), _obra(fila[8]),
                 _cuenta(fila[9]), _cuenta(fila[10]), (str(fila[11]).strip().upper() if fila[11] else None)),
            )
            pid = cur.lastrowid
            personas += 1
            for tipo, col in COLUMNAS.items():
                v = fila[col]
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                monto, estado, notas = None, "PENDIENTE", None
                if isinstance(v, (int, float)):
                    monto = round(float(v), 2)
                    estado = "PAGADO" if estado_total == "PAGADO" else "PENDIENTE"
                else:
                    t = v.strip().upper()
                    if t == "PAGADO":
                        estado = "PAGADO"
                    elif t != "NO PAGADO":
                        notas = v.strip()
                fecha_pago = None
                if tipo == "LIQUIDACION" and _fecha(fila[12]):
                    estado, fecha_pago = "PAGADO", _fecha(fila[12])
                    notas = "Fecha de comprobante tomada del Excel (falta subir el archivo)"
                c.execute(
                    "INSERT INTO items (personal_id, tipo, monto, estado, fecha_pago, notas) VALUES (?,?,?,?,?,?)",
                    (pid, tipo, monto, estado, fecha_pago, notas),
                )
                items += 1
    return personas, items


if __name__ == "__main__":
    ruta = sys.argv[1] if len(sys.argv) > 1 else "datos_iniciales.xlsx"
    p, i = importar(ruta)
    print(f"Importadas {p} personas y {i} conceptos.")
